import io
from collections.abc import Iterable
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1E293B")


def safe_cell(value: Any) -> Any:
    # Neutralize spreadsheet formula injection from scan content
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    if isinstance(value, str) and len(value) > 32000:
        return value[:32000] + "…"
    return value


def add_sheet(wb: Workbook, title: str, rows: Iterable[list[Any]]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.freeze_panes = "A2"
    first = True
    for row in rows:
        if first:
            cells = []
            for v in row:
                cell = WriteOnlyCell(ws, value=v)
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL
                cells.append(cell)
            ws.append(cells)
            first = False
        else:
            ws.append([safe_cell(v) for v in row])


def rows_to_xlsx(title: str, rows: Iterable[list[Any]]) -> bytes:
    return sheets_to_xlsx({title: rows})


def sheets_to_xlsx(sheets: dict[str, Iterable[list[Any]]]) -> bytes:
    wb = Workbook(write_only=True)
    for title, rows in sheets.items():
        add_sheet(wb, title, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
